import os
import sys
from unittest import mock

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
import asset_mapper as ed  # noqa: E402


# ---------- normalize / input
@pytest.mark.parametrize("raw,root,expected", [
    ("*.Dev.Example.com", "example.com", "dev.example.com"),
    ("www.example.com.", "example.com", "www.example.com"),
    ("example.com", "example.com", "example.com"),
    ("evil-example.com", "example.com", None),        # suffix-confusion
    ("example.com.evil.net", "example.com", None),
    ("other.org", "example.com", None),
    ("bad host.example.com", "example.com", None),
    ("a_b.example.com", "example.com", "a_b.example.com"),  # underscores are real in DNS
    ("-x.example.com", "example.com", None),
    ("user@example.com", "example.com", None),
    ("", "example.com", None),
])
def test_normalize(raw, root, expected):
    assert ed.normalize_host(raw, root) == expected


def test_load_domains(tmp_path):
    f = tmp_path / "d.txt"
    f.write_text("Example.com\nhttps://foo.org/path\n# comment\nexample.com\nnot a domain\n\nbar.io # trailing\n")
    assert ed.load_domains(str(f)) == ["example.com", "foo.org", "bar.io"]


def test_xl_injection():
    assert ed.xl("=cmd|' /C calc'!A0").startswith("'")
    assert ed.xl("@SUM(1)").startswith("'")
    assert ed.xl("normal.example.com") == "normal.example.com"
    assert ed.xl(5) == 5


# ---------- sources (mocked HTTP)
def fake_resp(status=200, payload=None):
    m = mock.Mock(status_code=status, text="")
    m.json.return_value = payload
    return m


def test_crtsh_parses_and_filters():
    rows = [{"name_value": "a.example.com\n*.b.example.com\nexample.com"},
            {"name_value": "evil.com\nexample.com.evil.com\nuser@example.com"}]
    with mock.patch.object(ed.requests, "get", return_value=fake_resp(200, rows)):
        hosts, err = ed.src_crtsh("example.com", {})
    assert err is None
    assert hosts == {"a.example.com", "b.example.com", "example.com"}


def test_crtsh_retries_then_reports_error():
    with mock.patch.object(ed.requests, "get", return_value=fake_resp(502)), \
         mock.patch.object(ed.time, "sleep"):
        hosts, err = ed.src_crtsh("example.com", {})
    assert hosts == set() and err == "HTTP 502"


def test_crtsh_recovers_after_502():
    seq = [fake_resp(502), fake_resp(200, [{"name_value": "x.example.com"}])]
    with mock.patch.object(ed.requests, "get", side_effect=seq), mock.patch.object(ed.time, "sleep"):
        hosts, err = ed.src_crtsh("example.com", {})
    assert hosts == {"x.example.com"} and err is None


def test_dnsdumpster_parse_and_headers():
    payload = {"a": [{"host": "www.example.com", "ips": [{"ip": "1.2.3.4"}]},
                     {"host": "evil.com", "ips": []}], "mx": [{"host": "mail.example.com"}]}
    with mock.patch.object(ed.requests, "get", return_value=fake_resp(200, payload)) as g, \
         mock.patch.object(ed.time, "sleep"):
        hosts, err = ed.src_dnsdumpster("example.com", {"dnsdumpster_key": "K"})
    assert hosts == {"www.example.com"} and err is None
    assert g.call_args.kwargs["headers"] == {"X-API-Key": "K"}


def test_dnsdumpster_no_key_and_auth_fail():
    assert ed.src_dnsdumpster("example.com", {})[1].startswith("skipped")
    with mock.patch.object(ed.requests, "get", return_value=fake_resp(401)), mock.patch.object(ed.time, "sleep"):
        assert "auth failed" in ed.src_dnsdumpster("example.com", {"dnsdumpster_key": "bad"})[1]


def test_subfinder_missing():
    with mock.patch.object(ed.shutil, "which", return_value=None):
        assert ed.src_subfinder("example.com", {})[1].startswith("skipped")


def test_subfinder_output_filtered():
    p = mock.Mock(returncode=0, stdout="a.example.com\nrandom.net\nB.Example.com\n", stderr="")
    with mock.patch.object(ed.shutil, "which", return_value="/x/subfinder"), \
         mock.patch.object(ed.subprocess, "run", return_value=p):
        hosts, err = ed.src_subfinder("example.com", {})
    assert hosts == {"a.example.com", "b.example.com"} and err is None


# ---------- DNS / ASN (live)
@pytest.mark.skipif(os.environ.get("SKIP_LIVE") == "1", reason="needs network (SKIP_LIVE=1)")
def test_live_resolve_and_asn():
    res = ed.make_resolver(["1.1.1.1", "8.8.8.8"])
    ips, status = ed.resolve_a(res, "example.com")
    assert status == "OK" and ips
    assert ed.resolve_a(res, "qa-does-not-exist-12345.invalid")[1] == "NXDOMAIN"
    info = ed.asn_lookup(res, "8.8.8.8", {})
    assert info["asn"] == "15169" and "GOOGLE" in info["name"].upper()
    assert ed.is_cloud(info["name"])


def test_asn_skips_private_ip():
    res = ed.make_resolver(["1.1.1.1"])
    assert ed.asn_lookup(res, "10.0.0.5", {})["asn"] == ""


# ---------- end to end (mocked sources, mocked resolution) -> xlsx
def test_end_to_end_workbook(tmp_path, monkeypatch):
    dom = tmp_path / "d.txt"
    dom.write_text("example.com\n")
    out = tmp_path / "out.xlsx"

    monkeypatch.setitem(ed.SOURCES, "crtsh", lambda d, c: ({"www.example.com", "=evil.example.com", "dead.example.com", "wild.example.com"}, None))
    monkeypatch.setitem(ed.SOURCES, "subfinder", lambda d, c: ({"www.example.com", "intra.example.com"}, None))
    monkeypatch.setitem(ed.SOURCES, "dnsdumpster", lambda d, c: (set(), "rate limited (429)"))

    table = {"example.com": (["93.184.216.34"], "OK"), "www.example.com": (["93.184.216.34", "8.8.8.8"], "OK"),
             "=evil.example.com": (["8.8.4.4"], "OK"),
             "dead.example.com": ([], "NXDOMAIN"), "intra.example.com": (["10.1.1.1"], "OK"),
             "wild.example.com": (["9.9.9.9"], "OK")}
    monkeypatch.setattr(ed, "resolve_a", lambda res, h: table.get(h, ([], "NXDOMAIN")) if not h.endswith(".example.com") or h in table else (["9.9.9.9"], "OK"))
    monkeypatch.setattr(ed, "asn_lookup", lambda res, ip, cache: cache.setdefault(
        ip, {"asn": "15169", "name": "GOOGLE, US", "prefix": "8.8.8.0/24", "country": "US"} if ip.startswith("8.") else
        {"asn": "1", "name": "SOMEISP, US", "prefix": "x/24", "country": "US"}))
    monkeypatch.setattr(sys, "argv", ["x", str(dom), "-o", str(out)])
    ed.main()

    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "All Records", "Unique IPs", "Unresolved", "Source Status"]
    ar = wb["All Records"]
    hdr = [c.value for c in ar[1]]
    assert hdr[0] == "Scope Validation Status" and hdr[-1] == "Notes"
    assert hdr[1:3] == ["Root Domain", "Hostname"] and "In Scope? (Client)" not in hdr
    rec = [[c.value for c in r] for r in ar.iter_rows(min_row=2)]
    for r in range(1, ar.max_row + 1):                     # light orange on first and last column
        assert ar[f"A{r}"].fill.fgColor.rgb.endswith("FCD5B4")
        assert ar[f"M{r}"].fill.fgColor.rgb.endswith("FCD5B4")
    assert ar["B2"].fill.fgColor.rgb != ar["A2"].fill.fgColor.rgb
    assert all(r[0] in (None, "") and r[-1] in (None, "") for r in rec)  # blank for the client
    dv = ar.data_validations.dataValidation
    assert len(dv) == 1 and "A2" in str(dv[0].sqref) and "Out of Scope" in dv[0].formula1
    hosts = {r[2] for r in rec}
    assert "dead.example.com" not in hosts
    assert "'=evil.example.com" in hosts                    # formula neutralised
    www = [r for r in rec if r[2] == "www.example.com"]
    assert len(www) == 2 and www[0][4] == "crtsh, subfinder" and www[0][5] == 2
    assert any(r[6] == "Yes" for r in www)                 # Google IP flagged cloud
    intra = [r for r in rec if r[2] == "intra.example.com"][0]
    assert "NON-PUBLIC IP" in intra[11]
    wild = [r for r in rec if r[2] == "wild.example.com"][0]
    assert "WILDCARD" in wild[11]                          # mock resolves any unknown host to 9.9.9.9
    apex = [r for r in rec if r[2] == "example.com"][0]
    assert apex[4] == "apex"
    unres = [[c.value for c in r] for r in wb["Unresolved"].iter_rows(min_row=2)]
    assert unres and unres[0][1] == "dead.example.com" and unres[0][3] == "NXDOMAIN"
    status = {(r[1].value): r[3].value for r in wb["Source Status"].iter_rows(min_row=2)}
    assert status["dnsdumpster"] == "rate limited (429)" and status["crtsh"] == "OK"
    ipsheet = [[c.value for c in r] for r in wb["Unique IPs"].iter_rows(min_row=2)]
    assert [r[0] for r in ipsheet] == sorted([r[0] for r in ipsheet], key=ed.ip_sort_key)


def test_end_to_end_all_sources_down(tmp_path, monkeypatch):
    dom = tmp_path / "d.txt"
    dom.write_text("example.com\n")
    out = tmp_path / "out.xlsx"
    for s in ed.SOURCES:
        monkeypatch.setitem(ed.SOURCES, s, lambda d, c: (set(), "HTTP 502"))
    monkeypatch.setattr(ed, "resolve_a", lambda r, h: ([], "NXDOMAIN"))
    monkeypatch.setattr(sys, "argv", ["x", str(dom), "-o", str(out)])
    ed.main()                                           # must not crash with empty data
    wb = load_workbook(out)
    assert wb["All Records"].max_row == 1


def test_bad_source_arg(tmp_path, monkeypatch):
    dom = tmp_path / "d.txt"
    dom.write_text("example.com\n")
    monkeypatch.setattr(sys, "argv", ["x", str(dom), "--sources", "bogus"])
    with pytest.raises(SystemExit):
        ed.main()


# ---------- subfinder discovery / diagnostics (fake binary)
def make_fake(tmp_path, body, name="subfinder"):
    p = tmp_path / name
    p.write_text("#!/bin/sh\n" + body + "\n")
    p.chmod(0o755)
    return str(p)


def test_subfinder_explicit_path_and_dedupe(tmp_path):
    exe = make_fake(tmp_path, 'printf "a.example.com\\nA.example.com\\nb.example.com\\nx.other.net\\n"')
    hosts, err = ed.src_subfinder("example.com", {"subfinder_path": exe})
    assert hosts == {"a.example.com", "b.example.com"} and err is None


def test_subfinder_found_outside_path(tmp_path, monkeypatch):
    gobin = tmp_path / "go" / "bin"
    gobin.mkdir(parents=True)
    make_fake(gobin, 'echo a.example.com', "subfinder")
    monkeypatch.setenv("PATH", "/usr/bin:/bin")
    monkeypatch.setattr(ed.os.path, "expanduser", lambda p: p.replace("~", str(tmp_path)))
    hosts, err = ed.src_subfinder("example.com", {})
    assert hosts == {"a.example.com"} and err is None


def test_subfinder_zero_results_is_loud(tmp_path):
    exe = make_fake(tmp_path, 'echo "[ERR] provider failed" >&2')
    hosts, err = ed.src_subfinder("example.com", {"subfinder_path": exe})
    assert hosts == set() and "0 usable results" in err and "provider failed" in err


def test_subfinder_nonzero_exit(tmp_path):
    exe = make_fake(tmp_path, 'echo boom >&2; exit 3')
    hosts, err = ed.src_subfinder("example.com", {"subfinder_path": exe})
    assert err.startswith("exit 3") and "boom" in err


def test_subfinder_bad_explicit_path():
    assert ed.src_subfinder("example.com", {"subfinder_path": "/nope/subfinder"})[1].startswith("skipped")
