#!/usr/bin/env python3
"""
External host discovery -> Excel.

Collects subdomains from several passive sources, resolves every hostname
yourself (A records only), enriches IPs with ASN data, and writes an .xlsx the
client can use to validate which external assets are in scope.

Sources: DNSDumpster API, crt.sh (certificate transparency), subfinder (if installed).

Setup:
    pip install requests openpyxl dnspython
    export DNSDUMPSTER_API_KEY=...        # optional; source skipped if absent

Usage:
    python3 asset_mapper.py domains.txt -o asset_mapper.xlsx
    python3 asset_mapper.py domains.txt --sources crtsh,subfinder

Only run against domains you are authorized to assess.
"""
import argparse
import ipaddress
import os
import random
import re
import shutil
import string
import subprocess
import sys
import time
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone

import dns.exception
import dns.resolver
import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DNSDUMPSTER_API = "https://api.dnsdumpster.com/domain/{}"
DNSDUMPSTER_DELAY = 2.1  # API allows 1 request / 2 seconds
LABEL = r"(?!-)[a-z0-9_-]{1,63}(?<!-)"
DOMAIN_RE = re.compile(rf"^{LABEL}(\.{LABEL})+$")
CLOUD_KEYWORDS = ("CLOUDFLARE", "AMAZON", "AWS", "AKAMAI", "FASTLY", "MICROSOFT", "AZURE",
                  "GOOGLE", "INCAPSULA", "IMPERVA", "SUCURI", "STACKPATH", "EDGECAST",
                  "DIGITALOCEAN", "LINODE", "OVH", "HETZNER", "ORACLE", "VERCEL", "NETLIFY",
                  "GITHUB", "HEROKU", "SALESFORCE", "ZSCALER")

CLIENT_FILL = "FCD5B4"  # light orange: cells the client fills in
SCOPE_OPTIONS = "In Scope,Out of Scope,Unsure"
HOST_COLS = ["Scope Validation Status", "Root Domain", "Hostname", "IP Address", "Sources",
             "# Sources", "Cloud/CDN", "ASN", "ASN Name", "Network Range", "Country", "Flags",
             "Notes"]


# ---------------------------------------------------------------- input / helpers

def normalize_host(h, root=None):
    """Lowercase, strip wildcard/trailing dot; return None if invalid or outside root."""
    h = h.strip().lower().lstrip("*.").rstrip(".")
    if not DOMAIN_RE.match(h):
        return None
    if root and h != root and not h.endswith("." + root):
        return None
    return h


def load_domains(path):
    domains, seen = [], set()
    with open(path) as f:
        for line in f:
            d = line.split("#")[0].strip().lower()
            d = re.sub(r"^https?://", "", d).split("/")[0].rstrip(".")
            if not d:
                continue
            if not DOMAIN_RE.match(d):
                print(f"[!] Skipping invalid domain: {d}", file=sys.stderr)
                continue
            if d not in seen:
                seen.add(d)
                domains.append(d)
    return domains


def xl(v):
    """Neutralize spreadsheet formula injection from third-party data."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


# ---------------------------------------------------------------- sources
# Each returns (set_of_hostnames, error_message_or_None)

def src_dnsdumpster(domain, ctx):
    key = ctx.get("dnsdumpster_key")
    if not key:
        return set(), "skipped (no DNSDUMPSTER_API_KEY)"
    elapsed = time.monotonic() - ctx.get("dd_last", 0)
    if elapsed < DNSDUMPSTER_DELAY:
        time.sleep(DNSDUMPSTER_DELAY - elapsed)
    err = "failed"
    for attempt in range(3):
        try:
            r = requests.get(DNSDUMPSTER_API.format(domain), headers={"X-API-Key": key}, timeout=60)
            ctx["dd_last"] = time.monotonic()
        except requests.RequestException as e:
            err = str(e)[:100]
            time.sleep(5)
            continue
        if r.status_code == 200:
            hosts = {normalize_host(rec.get("host", ""), domain) for rec in (r.json().get("a") or [])}
            return {h for h in hosts if h}, None
        if r.status_code in (401, 403):
            return set(), f"auth failed (HTTP {r.status_code})"
        if r.status_code == 429:
            err = "rate limited (429)"
            time.sleep(5 * (attempt + 1))
            continue
        return set(), f"HTTP {r.status_code}"
    return set(), err


def src_crtsh(domain, ctx):
    err = "failed"
    for attempt in range(4):
        try:
            r = requests.get("https://crt.sh/", params={"q": f"%.{domain}", "output": "json"}, timeout=90)
        except requests.RequestException as e:
            err = str(e)[:100]
            time.sleep(4 * (attempt + 1))
            continue
        if r.status_code == 200:
            try:
                rows = r.json()
            except ValueError:
                err = "invalid JSON"
                time.sleep(4 * (attempt + 1))
                continue
            hosts = set()
            for row in rows:
                for name in str(row.get("name_value", "")).splitlines():
                    h = normalize_host(name, domain)
                    if h:
                        hosts.add(h)
            return hosts, None
        err = f"HTTP {r.status_code}"
        time.sleep(4 * (attempt + 1))
    return set(), err


def find_subfinder(explicit=None):
    """PATH lookup first, then common install dirs (Go/Homebrew) that GUI/IDE shells often miss."""
    if explicit:
        return explicit if os.access(explicit, os.X_OK) else None
    found = shutil.which("subfinder")
    if found:
        return found
    for d in ("~/go/bin", "/opt/homebrew/bin", "/usr/local/bin", "/usr/local/go/bin", "~/.local/bin"):
        p = os.path.join(os.path.expanduser(d), "subfinder")
        if os.access(p, os.X_OK):
            return p
    return None


def src_subfinder(domain, ctx):
    exe = find_subfinder(ctx.get("subfinder_path"))
    if not exe:
        return set(), "skipped (subfinder not found; use --subfinder-path)"
    try:
        p = subprocess.run([exe, "-d", domain, "-silent"], capture_output=True, text=True, timeout=ctx.get("subfinder_timeout", 900))
    except subprocess.TimeoutExpired:
        return set(), "timed out"
    if p.returncode != 0:
        return set(), f"exit {p.returncode}: {p.stderr.strip()[:150]}"
    lines = [x.strip() for x in p.stdout.splitlines() if x.strip()]
    hosts = {h for h in (normalize_host(re.split(r"[,\s]", x)[0], domain) for x in lines) if h}
    if not hosts:
        return set(), (f"0 usable results ({exe}; {len(lines)} raw lines; "
                       f"stderr: {p.stderr.strip()[:150] or 'none'})")
    if len(hosts) < len(lines):
        print(f"    subfinder: {len(lines)} raw lines -> {len(hosts)} unique in-scope hosts", file=sys.stderr)
    return hosts, None


SOURCES = {"dnsdumpster": src_dnsdumpster, "crtsh": src_crtsh, "subfinder": src_subfinder}


# ---------------------------------------------------------------- DNS resolution / ASN

def make_resolver(nameservers):
    r = dns.resolver.Resolver(configure=False)
    r.nameservers = nameservers
    r.lifetime = 6
    r.timeout = 3
    return r


def resolve_a(res, host):
    """Return (sorted_ips, status) where status in OK / NXDOMAIN / NOANSWER / ERROR."""
    try:
        ans = res.resolve(host, "A")
        return sorted({a.address for a in ans}), "OK"
    except dns.resolver.NXDOMAIN:
        return [], "NXDOMAIN"
    except dns.resolver.NoAnswer:
        return [], "NOANSWER"
    except (dns.exception.DNSException, OSError):
        return [], "ERROR"


def wildcard_ips(res, domain):
    probe = "".join(random.choices(string.ascii_lowercase, k=16)) + "." + domain
    ips, status = resolve_a(res, probe)
    return set(ips) if status == "OK" else set()


def asn_lookup(res, ip, cache):
    """Team Cymru DNS lookup. Returns dict(asn, name, prefix, country)."""
    if ip in cache:
        return cache[ip]
    info = {"asn": "", "name": "", "prefix": "", "country": ""}
    try:
        a = ipaddress.ip_address(ip)
        if a.version == 4 and a.is_global:
            q = ".".join(reversed(ip.split("."))) + ".origin.asn.cymru.com"
            txt = str(res.resolve(q, "TXT")[0]).strip('"')
            asn, prefix, country = [p.strip() for p in txt.split("|")[:3]]
            asn = asn.split()[0]
            info.update(asn=asn, prefix=prefix, country=country)
            try:
                t2 = str(res.resolve(f"AS{asn}.asn.cymru.com", "TXT")[0]).strip('"')
                info["name"] = t2.split("|")[-1].strip()
            except dns.exception.DNSException:
                pass
    except (dns.exception.DNSException, ValueError, IndexError, OSError):
        pass
    cache[ip] = info
    return info


def is_cloud(name):
    n = name.upper()
    return any(k in n for k in CLOUD_KEYWORDS)


# ---------------------------------------------------------------- Excel

def style_sheet(ws):
    fill = PatternFill("solid", fgColor="1F3864")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.columns, 1):
        longest = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 12), 60)


def add_scope_dropdown(ws, col_letter, options="Yes,No,Unsure"):
    if ws.max_row < 2:
        return
    dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{ws.max_row}")


def highlight_client_column(ws, col_letter, width=None):
    """Light-orange fill on a column the client must fill in (header + every row)."""
    fill = PatternFill("solid", fgColor=CLIENT_FILL)
    ws[f"{col_letter}1"].font = Font(bold=True, color="000000")
    for row in range(1, ws.max_row + 1):
        ws[f"{col_letter}{row}"].fill = fill
    if width:
        ws.column_dimensions[col_letter].width = width


def ip_sort_key(ip):
    try:
        return (0, int(ipaddress.ip_address(ip)))
    except ValueError:
        return (1, 0)


def write_workbook(path, domains, results, source_status, resolved_rows, unresolved_rows, source_names):
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws.append(["Root Domain"] + [f"Found by {s}" for s in source_names] +
              ["Total Unique Hosts", "Resolved", "Unresolved", "Unique IPs"])
    for d in domains:
        hosts = results[d]
        rrows = [r for r in resolved_rows if r["domain"] == d]
        ws.append([d] + [sum(1 for h in hosts.values() if s in h) for s in source_names] +
                  [len(hosts), len({r["host"] for r in rrows}),
                   len([r for r in unresolved_rows if r["domain"] == d]),
                   len({r["ip"] for r in rrows})])
    ws.append([])
    ws.append([f"Generated {datetime.now(timezone.utc):%Y-%m-%d %H:%M UTC}. Sources are passive; "
               f"resolution done live via public resolvers."])
    style_sheet(ws)

    ws = wb.create_sheet("All Records")
    ws.append(HOST_COLS)
    for r in sorted(resolved_rows, key=lambda r: (r["domain"], r["host"], ip_sort_key(r["ip"]))):
        ws.append([xl(v) for v in ["", r["domain"], r["host"], r["ip"], ", ".join(r["sources"]),
                                   len(r["sources"]), "Yes" if r["cloud"] else "No",
                                   r["asn"], r["asn_name"], r["prefix"], r["country"],
                                   ", ".join(r["flags"]), ""]])
    style_sheet(ws)
    last = get_column_letter(len(HOST_COLS))
    add_scope_dropdown(ws, "A", SCOPE_OPTIONS)
    highlight_client_column(ws, "A", width=26)
    highlight_client_column(ws, last, width=40)

    by_ip = defaultdict(lambda: {"hosts": set(), "row": None})
    for r in resolved_rows:
        e = by_ip[r["ip"]]
        e["hosts"].add(r["host"])
        e["row"] = e["row"] or r
    ws = wb.create_sheet("Unique IPs")
    ws.append(["IP Address", "Hostnames", "Host Count", "Cloud/CDN", "ASN", "ASN Name",
               "Network Range", "Country", "Flags", "In Scope? (Client)", "Client Notes"])
    for ip in sorted(by_ip, key=ip_sort_key):
        e, r = by_ip[ip], by_ip[ip]["row"]
        ws.append([xl(v) for v in [ip, "\n".join(sorted(e["hosts"])), len(e["hosts"]),
                                   "Yes" if r["cloud"] else "No", r["asn"], r["asn_name"],
                                   r["prefix"], r["country"], ", ".join(r["flags"]), "", ""]])
    style_sheet(ws)
    add_scope_dropdown(ws, "J")

    ws = wb.create_sheet("Unresolved")
    ws.append(["Root Domain", "Hostname", "Sources", "DNS Status"])
    for r in sorted(unresolved_rows, key=lambda r: (r["domain"], r["host"])):
        ws.append([xl(v) for v in [r["domain"], r["host"], ", ".join(r["sources"]), r["status"]]])
    style_sheet(ws)

    ws = wb.create_sheet("Source Status")
    ws.append(["Root Domain", "Source", "Hosts Returned", "Status"])
    for (d, s), (n, err) in source_status.items():
        ws.append([d, s, n, err or "OK"])
    style_sheet(ws)

    wb.save(path)


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("domains_file", help="Text file, one domain per line (# comments allowed)")
    ap.add_argument("-o", "--output", default="asset_mapper.xlsx")
    ap.add_argument("--sources", default="dnsdumpster,crtsh,subfinder",
                    help="Comma-separated subset of: " + ",".join(SOURCES))
    ap.add_argument("-k", "--api-key", default=os.environ.get("DNSDUMPSTER_API_KEY"),
                    help="DNSDumpster API key (default: $DNSDUMPSTER_API_KEY)")
    ap.add_argument("--resolvers", default="1.1.1.1,8.8.8.8", help="Comma-separated DNS resolvers")
    ap.add_argument("--threads", type=int, default=20)
    ap.add_argument("--subfinder-path", help="Full path to the subfinder binary (if not on PATH)")
    ap.add_argument("--no-asn", action="store_true", help="Skip ASN/Cloud enrichment")
    args = ap.parse_args()

    source_names = [s.strip() for s in args.sources.split(",") if s.strip()]
    bad = [s for s in source_names if s not in SOURCES]
    if bad:
        sys.exit(f"[x] Unknown source(s): {', '.join(bad)}")

    try:
        domains = load_domains(args.domains_file)
    except OSError as e:
        sys.exit(f"[x] Cannot read {args.domains_file}: {e.strerror}")
    if not domains:
        sys.exit("[x] No valid domains found.")

    ctx = {"dnsdumpster_key": args.api_key, "subfinder_path": args.subfinder_path}
    results = {d: {} for d in domains}  # domain -> {host: set(sources)}
    source_status = {}

    # 1. collect
    for i, d in enumerate(domains, 1):
        print(f"[{i}/{len(domains)}] {d}")
        for s in source_names:
            hosts, err = SOURCES[s](d, ctx)
            source_status[(d, s)] = (len(hosts), err)
            for h in hosts:
                results[d].setdefault(h, set()).add(s)
            print(f"    {s}: {len(hosts)} hosts" + (f"  [{err}]" if err else ""))
        results[d].setdefault(d, set())  # always check the apex itself

    # 2. resolve
    res = make_resolver([r.strip() for r in args.resolvers.split(",") if r.strip()])
    wild = {d: wildcard_ips(res, d) for d in domains}
    for d, w in wild.items():
        if w:
            print(f"[!] {d} has wildcard DNS ({', '.join(sorted(w))}); affected hosts are flagged")

    jobs = [(d, h) for d in domains for h in results[d]]
    print(f"[*] Resolving {len(jobs)} hostnames...")
    with ThreadPoolExecutor(max_workers=args.threads) as ex:
        resolved = list(ex.map(lambda j: resolve_a(res, j[1]), jobs))

    # 3. enrich
    asn_cache = {}
    ips_all = sorted({ip for ips, _ in resolved for ip in ips})
    if not args.no_asn:
        print(f"[*] ASN lookup for {len(ips_all)} unique IPs...")
        with ThreadPoolExecutor(max_workers=args.threads) as ex:
            list(ex.map(lambda ip: asn_lookup(res, ip, asn_cache), ips_all))

    resolved_rows, unresolved_rows = [], []
    for (d, h), (ips, status) in zip(jobs, resolved):
        srcs = sorted(results[d][h]) or ["apex"]
        if not ips:
            unresolved_rows.append({"domain": d, "host": h, "sources": srcs, "status": status})
            continue
        for ip in ips:
            info = asn_cache.get(ip, {"asn": "", "name": "", "prefix": "", "country": ""})
            flags = []
            if not ipaddress.ip_address(ip).is_global:
                flags.append("NON-PUBLIC IP")
            if wild[d] and set(ips) <= wild[d]:
                flags.append("WILDCARD")
            if len(results[d][h]) == 1 and srcs != ["apex"]:
                flags.append("SINGLE-SOURCE")
            resolved_rows.append({"domain": d, "host": h, "ip": ip, "sources": srcs,
                                  "cloud": is_cloud(info["name"]), "asn": info["asn"],
                                  "asn_name": info["name"], "prefix": info["prefix"],
                                  "country": info["country"], "flags": flags})

    write_workbook(args.output, domains, results, source_status, resolved_rows,
                   unresolved_rows, source_names)
    print(f"\n[+] {len(resolved_rows)} resolved records, {len(unresolved_rows)} unresolved hosts -> {args.output}")
    failed = [(d, s, e) for (d, s), (n, e) in source_status.items() if e]
    if failed:
        print("[!] Some sources failed; coverage may be incomplete (see 'Source Status' sheet):")
        for d, s, e in failed:
            print(f"    {d} / {s}: {e}")


if __name__ == "__main__":
    main()
